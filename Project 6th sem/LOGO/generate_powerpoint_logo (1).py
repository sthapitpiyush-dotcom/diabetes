from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

def create_aimid_logo_presentation():
    # Initialize 16:9 widescreen presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Use a blank slide layout
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)

    # Color definitions (exact matching from image_ce0162.png)
    bg_color = RGBColor(11, 15, 25)          # Deep dark navy
    cyan_color = RGBColor(103, 232, 249)     # Bright cyan center pillar
    slate_color = RGBColor(55, 65, 81)       # Slate gray side pillars
    beam_color = RGBColor(31, 41, 55)        # Dark slate gray angled beams
    white_color = RGBColor(255, 255, 255)    # Title text
    tagline_color = RGBColor(156, 163, 175)  # Tagline text

    # 1. Add background shape to cover the entire slide
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = bg_color
    bg_shape.line.fill.background()  # No outline

    # --- VECTOR PILLARS DEFINITIONS ---
    # We define sizes and align bottoms at Y = 4.2 inches
    bottom_y = 4.2
    pillar_width = 0.45
    gap = 0.35
    center_x = 6.666  # Center of widescreen slide (13.333 / 2)

    # Center Cyan Pillar
    center_height = 2.8
    center_left = center_x - (pillar_width / 2)
    center_top = bottom_y - center_height

    center_pillar = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, 
        Inches(center_left), Inches(center_top), 
        Inches(pillar_width), Inches(center_height)
    )
    center_pillar.fill.solid()
    center_pillar.fill.fore_color.rgb = cyan_color
    center_pillar.line.color.rgb = cyan_color

    # Left Slate Pillar
    side_height = 2.0
    left_pillar_left = center_left - pillar_width - gap
    left_pillar_top = bottom_y - side_height

    left_pillar = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left_pillar_left), Inches(left_pillar_top),
        Inches(pillar_width), Inches(side_height)
    )
    left_pillar.fill.solid()
    left_pillar.fill.fore_color.rgb = slate_color
    left_pillar.line.color.rgb = slate_color

    # Right Slate Pillar
    right_pillar_left = center_left + pillar_width + gap
    right_pillar_top = bottom_y - side_height

    right_pillar = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(right_pillar_left), Inches(right_pillar_top),
        Inches(pillar_width), Inches(side_height)
    )
    right_pillar.fill.solid()
    right_pillar.fill.fore_color.rgb = slate_color
    right_pillar.line.color.rgb = slate_color

    # --- ANGLED BRIDGING BEAMS ---
    # Left Beam (sits over left pillar pointing up-right)
    left_beam = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(center_left - 0.55), Inches(center_top - 0.05),
        Inches(0.65), Inches(0.18)
    )
    left_beam.fill.solid()
    left_beam.fill.fore_color.rgb = beam_color
    left_beam.line.color.rgb = beam_color
    left_beam.rotation = -28  # Slanted upwards

    # Right Beam (sits between center pillar and right pillar pointing down-right)
    right_beam = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(center_left + 0.35), Inches(center_top + 0.25),
        Inches(0.7), Inches(0.18)
    )
    right_beam.fill.solid()
    right_beam.fill.fore_color.rgb = beam_color
    right_beam.line.color.rgb = beam_color
    right_beam.rotation = 32  # Slanted downwards

    # --- TYPOGRAPHY / TEXT FIELDS ---
    # Main Brand Text: "aimid"
    title_box = slide.shapes.add_textbox(Inches(3.66), Inches(4.5), Inches(6.0), Inches(1.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "aimid"
    p.font.name = "Segoe UI"  # Standard clean geometric font
    p.font.size = Pt(56)
    p.font.bold = True
    p.font.color.rgb = white_color
    p.alignment = PP_ALIGN.CENTER

    # Subtitle/Tagline Text: "AI MIDDLEWARE FOR THE ENTERPRISE"
    tag_box = slide.shapes.add_textbox(Inches(2.66), Inches(5.6), Inches(8.0), Inches(0.8))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = "AI MIDDLEWARE FOR THE ENTERPRISE"
    p_tag.font.name = "Segoe UI"
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = tagline_color
    # Emulate wide character tracking spacing using spaces
    p_tag.text = "A I   M I D D L E W A R E   F O R   T H E   E N T E R P R I S E"
    p_tag.alignment = PP_ALIGN.CENTER

    # Save presentation
    output_filename = "aimid_logo_run2.pptx"
    prs.save(output_filename)
    print(f"Successfully generated clean, editable vector slide: '{output_filename}'")

if __name__ == "__main__":
    create_aimid_logo_presentation()
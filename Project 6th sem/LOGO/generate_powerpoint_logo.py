import subprocess
import sys

# Automatically install python-pptx if it's not present
try:
    import pptx
except ImportError:
    print("Installing required python-pptx library...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    from pptx import Presentation

from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_editable_logo_pptx(filename="aimid_logo.pptx"):
    prs = Presentation()
    
    # Configure to 16:9 Widescreen (13.33 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Use blank slide layout
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    
    # 1. Background (Dark Obsidian)
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        Inches(0), Inches(0), 
        prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(11, 15, 25) # #0B0F19
    bg.line.fill.background() # No border
    
    # Coordinates calculation for centering on 13.33" x 7.5"
    # Center X: 6.66", Center Y: 3.75"
    
    # 2. Central Pillar (Kinetic Teal)
    center_pillar = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.46), Inches(2.5), Inches(0.4), Inches(2.5)
    )
    center_pillar.fill.solid()
    center_pillar.fill.fore_color.rgb = RGBColor(0, 242, 254) # #00F2FE
    center_pillar.line.fill.background()
    
    # 3. Left Pillar (Secure Slate)
    left_pillar = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(5.56), Inches(3.125), Inches(0.4), Inches(1.875)
    )
    left_pillar.fill.solid()
    left_pillar.fill.fore_color.rgb = RGBColor(51, 65, 85) # #334155
    left_pillar.line.fill.background()
    
    # 4. Right Pillar (Secure Slate)
    right_pillar = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.36), Inches(3.125), Inches(0.4), Inches(1.875)
    )
    right_pillar.fill.solid()
    right_pillar.fill.fore_color.rgb = RGBColor(51, 65, 85) # #334155
    right_pillar.line.fill.background()
    
    # 5. Left Bridge (Angled middleware path)
    left_bridge = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(5.82), Inches(2.85), Inches(0.7), Inches(0.2)
    )
    left_bridge.fill.solid()
    left_bridge.fill.fore_color.rgb = RGBColor(30, 41, 59) # #1E293B
    left_bridge.line.fill.background()
    left_bridge.rotation = -35  # Angled upward to central pillar
    
    # 6. Right Bridge (Angled middleware path)
    right_bridge = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(6.82), Inches(2.85), Inches(0.7), Inches(0.2)
    )
    right_bridge.fill.solid()
    right_bridge.fill.fore_color.rgb = RGBColor(30, 41, 59) # #1E293B
    right_bridge.line.fill.background()
    right_bridge.rotation = 35   # Angled upward from right pillar
    
    # Save Presentation
    prs.save(filename)
    print(f"Success! '{filename}' has been created with fully editable vector shapes.")

if __name__ == "__main__":
    create_editable_logo_pptx()
